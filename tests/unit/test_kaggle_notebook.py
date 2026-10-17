"""Notebook behavior without Kaggle credentials, downloads, or GPU allocation."""
import ast
import json
import os
from pathlib import Path
import subprocess
import sys
import types
from zipfile import ZipFile

import pymupdf
import pytest

ROOT = Path(__file__).resolve().parents[2]
NOTEBOOK = ROOT / "kaggle/ocr_pdf_kaggle.ipynb"


def _cells():
    notebook = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    return notebook, {cell["id"]: "".join(cell["source"]) for cell in notebook["cells"]
                      if cell["cell_type"] == "code"}


def _namespace():
    _, cells = _cells()
    namespace = {}
    for name in ("parameters", "helpers"):
        exec(compile(cells[name], str(NOTEBOOK) + ":" + name, "exec"), namespace)
    return namespace, cells


def _functions(source, namespace):
    tree = ast.parse(source)
    tree.body = [node for node in tree.body if isinstance(node, ast.FunctionDef)]
    exec(compile(tree, str(NOTEBOOK), "exec"), namespace)


def test_notebook_is_clean_and_all_code_is_python():
    notebook, cells = _cells()
    assert notebook["nbformat"] == 4
    assert notebook["metadata"]["kaggle"]["isInternetEnabled"] is True
    assert len(cells) == 9
    for source in cells.values():
        ast.parse(source)
    assert all(not cell["outputs"] and cell["execution_count"] is None
               for cell in notebook["cells"] if cell["cell_type"] == "code")
    namespace, _ = _namespace()
    assert namespace["FILE_LIMIT"] is None
    assert namespace["PDF_FILES"] == []
    assert namespace["RUN_FULL_BATCH"] is True


@pytest.mark.parametrize("exit_code", [0, 7])
def test_command_log_preserves_output_and_propagates_failure(tmp_path, exit_code):
    namespace, _ = _namespace()
    log = tmp_path / "command.log"
    command = [sys.executable, "-c", f"print('notebook-log'); raise SystemExit({exit_code})"]
    if exit_code:
        with pytest.raises(subprocess.CalledProcessError) as error:
            namespace["run_logged"](command, log, echo=False)
        assert error.value.returncode == exit_code
    else:
        namespace["run_logged"](command, log, echo=False)
    assert "notebook-log" in log.read_text(encoding="utf-8")


def test_input_selection_handles_uppercase_extensions_and_duplicate_names(tmp_path):
    namespace, cells = _namespace()
    _functions(cells["inputs"], namespace)
    inputs = tmp_path / "inputs"
    originals = []
    for folder, name in (("first", "annual.pdf"), ("second", "annual.PDF")):
        path = inputs / folder / name
        path.parent.mkdir(parents=True)
        path.write_bytes(b"pdf fixture")
        originals.append(path)
    paths = namespace["select_pdfs"](inputs)
    staged, mapping = namespace["stage_pdfs"](paths, tmp_path / "runtime")
    assert len(mapping) == 2
    assert len({item["document_id"].casefold() for item in mapping}) == 2
    assert len(list(staged.glob("*.pdf"))) == 2
    assert all(path.read_bytes() == b"pdf fixture" for path in originals)
    assert len(namespace["select_pdfs"](inputs, limit=1)) == 1
    assert namespace["select_pdfs"](inputs, ["second/annual.PDF"]) == [originals[1].resolve()]
    with pytest.raises(ValueError):
        namespace["select_pdfs"](inputs, limit=0)
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"outside")
    with pytest.raises(ValueError):
        namespace["select_pdfs"](inputs, [str(outside)])


def test_missing_pdfs_fail_before_installing_anything(tmp_path):
    namespace, cells = _namespace()
    namespace.update(INPUT_ROOT=tmp_path, RUNTIME_ROOT=tmp_path / "runtime", INPUTS_READY=True)
    with pytest.raises(FileNotFoundError, match="Add Input"):
        exec(cells["inputs"], namespace)
    assert not (tmp_path / "runtime").exists()
    assert namespace["INPUTS_READY"] is False


def test_setup_requires_successful_input_cell_and_config_rerun_invalidates_it(tmp_path):
    namespace, cells = _namespace()
    namespace["run_logged"] = lambda *a, **kw: pytest.fail("Must not install before input selection")
    with pytest.raises(RuntimeError, match="cell 3"):
        exec(cells["setup"], namespace)
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    (inputs / "report.pdf").write_bytes(b"fixture")
    results = tmp_path / "results"
    namespace.update(INPUT_ROOT=inputs, RUNTIME_ROOT=tmp_path / "runtime",
                     RESULTS_DIR=results, LOG_DIR=results / "logs")
    exec(cells["inputs"], namespace)
    assert namespace["INPUTS_READY"] is True
    exec(cells["parameters"], namespace)
    assert namespace["INPUTS_READY"] is False
    with pytest.raises(RuntimeError, match="Run All"):
        exec(cells["setup"], namespace)


def test_batch_checks_file_status_instead_of_only_process_exit_code(tmp_path):
    namespace, cells = _namespace()
    exec(cells["commands"], namespace)
    manifest = tmp_path / "batch_manifest.json"
    manifest.write_text(json.dumps([{"document_id": "annual", "status": "FAILED_PROCESSING",
                                     "error": "OCR failure"}]), encoding="utf-8")
    with pytest.raises(RuntimeError, match="OCR failure"):
        namespace["check_batch_manifest"](tmp_path, ["annual"])
    manifest.write_text('[{"document_id":"annual","status":"CACHED"}]', encoding="utf-8")
    assert len(namespace["check_batch_manifest"](tmp_path, ["annual"])) == 1
    with pytest.raises(RuntimeError, match="PDF"):
        namespace["check_batch_manifest"](tmp_path, ["missing"])


@pytest.mark.parametrize("full_batch", [False, True])
def test_native_notebook_flow_creates_excel_and_download_zip(tmp_path, full_batch):
    namespace, cells = _namespace()
    inputs = tmp_path / "input"
    inputs.mkdir()
    with pymupdf.open() as source:
        page = source.new_page()
        page.insert_textbox((40, 40, 550, 700), (
            "Annual financial report for shareholders and investors. "
            "The company uses artificial intelligence and cloud computing. "
            "We invest in artificial intelligence for business growth. "
            "Management develops technology to improve customer service and operations. "
            "Revenue and profit are reviewed by the board each year.\n") * 4)
        source.save(inputs / "annual.pdf")
    results = tmp_path / "results"
    runtime = tmp_path / "runtime"
    namespace.update(PROFILE="fast", DEVICE="cpu", INPUT_ROOT=inputs,
        RUNTIME_ROOT=runtime, RESULTS_DIR=results, LOG_DIR=results / "logs",
        SMOKE_DIR=results / "smoke", BATCH_DIR=results / "documents",
        ARCHIVE_PATH=tmp_path / "results.zip", RUN_FULL_BATCH=full_batch, ASSETS="none")
    exec(cells["inputs"], namespace)
    keywords = tmp_path / "keywords.txt"
    keywords.write_text("artificial intelligence\nabsent\n", encoding="utf-8")
    hardware = runtime / "hardware.yaml"
    hardware.write_text(json.dumps({"hardware": {"name": "notebook-test", "device": "cpu",
        "workers": 1, "threads_per_worker": 1, "batch_size": 4}}), encoding="utf-8")
    environment = os.environ.copy()
    environment.pop("OCR_CONFIG_PATH", None)
    environment.update(PYTHONIOENCODING="utf-8", PYTHONPATH=str(ROOT),
        OCR_OUTPUT_DIR=str(results / "documents"), OCR_CACHE_DIR=str(runtime / "cache"))
    namespace.update(REPO_DIR=ROOT, OCR_PYTHON=Path(sys.executable), CACHE_DIR=runtime / "cache",
        KEYWORDS_PATH=keywords, HARDWARE_PATH=hardware, RUNTIME_ENV=environment)
    for name in ("commands", "smoke", "batch", "matrix"):
        exec(cells[name], namespace)
    assert namespace["SMOKE_SUMMARY"][0]["matches"] == 8
    if full_batch:
        assert namespace["BATCH_SUMMARY"][0]["matches"] == 8
        assert namespace["MATRIX_PATH"].name == "MA_TRAN_DEM_TU_KHOA_1_FILES.xlsx"
    else:
        assert namespace["MATRIX_PATH"] is None
    _functions(cells["download"], namespace)
    archive = namespace["bundle_results"](results, tmp_path / "results.zip")
    with ZipFile(archive) as bundle:
        names = bundle.namelist()
        assert "smoke/annual/count_keywords.xlsx" in names
        assert ("documents/annual/count_keywords.xlsx" in names) == full_batch
        assert ("MA_TRAN_DEM_TU_KHOA_1_FILES.xlsx" in names) == full_batch
        assert not any(".venv" in name or name.startswith("models/") for name in names)


def _matrix_namespace(tmp_path, texts, *, completed=True):
    from src.ocr.infrastructure.export.completion import mark_completed
    namespace, cells = _namespace()
    results = tmp_path / "results"
    batch = results / "documents"
    batch.mkdir(parents=True)
    mapping = []
    for index, text in enumerate(texts):
        document_id = f"annual_{index + 1:02d}"
        directory = batch / document_id
        directory.mkdir()
        (directory / "result.md").write_text(text, encoding="utf-8")
        if completed:
            mark_completed(directory, "PASS")
        mapping.append({"document_id": document_id, "source": f"/input/{document_id}.pdf",
                        "staged": f"/staged/{document_id}.pdf"})
    (batch / "batch_manifest.json").write_text(json.dumps([
        {"document_id": item["document_id"], "status": "COMPLETED"} for item in mapping
    ]), encoding="utf-8")
    namespace.update(REPO_DIR=ROOT, OCR_PYTHON=Path(sys.executable),
        RESULTS_DIR=results, LOG_DIR=results / "logs", BATCH_DIR=batch,
        INPUT_MAPPING=mapping, RUN_FULL_BATCH=True,
        RUNTIME_ENV={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONPATH": str(ROOT)})
    exec(cells["commands"], namespace)
    return namespace, cells


@pytest.mark.parametrize("document_count", [50, 53])
def test_matrix_exports_every_file_and_exactly_88_ordered_keyword_columns(tmp_path, document_count):
    from openpyxl import load_workbook
    texts = ["Trí tuệ nhân tạo. " + "Internet di động. " * (index % 3) +
             "HỌC MÁY. " * (index % 2) for index in range(document_count)]
    namespace, cells = _matrix_namespace(tmp_path, texts)
    expected_keywords = [line.strip() for line in (ROOT / "config/keywords.txt").read_text(
        encoding="utf-8-sig").splitlines() if line.strip()]
    assert namespace["MATRIX_KEYWORDS"] == expected_keywords
    exec(cells["matrix"], namespace)
    assert namespace["MATRIX_PATH"].name == f"MA_TRAN_DEM_TU_KHOA_{document_count}_FILES.xlsx"
    workbook = load_workbook(namespace["MATRIX_PATH"], data_only=True)
    try:
        sheet = workbook["Ma trận từ khóa"]
        values = list(sheet.iter_rows(values_only=True))
        assert sheet.max_row == document_count + 1 and sheet.max_column == 89
        assert list(values[0]) == ["Tệp PDF", *expected_keywords]
        assert values[0][43] == values[0][45] == "Internet di động"
        assert sheet.freeze_panes == "B2"
        assert sheet.auto_filter.ref == f"A1:CK{document_count + 1}"
        for index, row in enumerate(values[1:]):
            assert row[0] == f"annual_{index + 1:02d}.pdf"
            assert row[1] == 1  # Trí tuệ nhân tạo.
            assert row[7] == index % 2  # học máy, ignoring case.
            assert row[43] == row[45] == index % 3
            assert row[6] == 0  # robot thông minh is absent.
            assert all(type(count) is int for count in row[1:])
        info = workbook["Thông tin"]
        assert info["B2"].value == document_count
        assert info["B3"].value == 88
        assert info["B4"].value == 87
    finally:
        workbook.close()


@pytest.mark.parametrize("tampered", [False, True])
def test_matrix_does_not_turn_incomplete_or_damaged_documents_into_zero_counts(tmp_path, tampered):
    namespace, cells = _matrix_namespace(tmp_path, ["Internet di động"], completed=tampered)
    if tampered:
        (namespace["BATCH_DIR"] / "annual_01/result.md").write_text("modified", encoding="utf-8")
    with pytest.raises(subprocess.CalledProcessError):
        exec(cells["matrix"], namespace)
    assert not list(namespace["RESULTS_DIR"].glob("MA_TRAN_DEM_TU_KHOA_*.xlsx"))


@pytest.mark.parametrize("cross_volume", [False, True])
def test_zip_can_be_downloaded_when_an_earlier_cell_failed_before_matrix_export(tmp_path, monkeypatch, cross_volume):
    namespace, cells = _namespace()
    results = tmp_path / "results"
    results.mkdir()
    (results / "error.log").write_text("Earlier cell failed", encoding="utf-8")
    namespace.update(RESULTS_DIR=results, ARCHIVE_PATH=tmp_path / "results.zip")
    if cross_volume:
        def other_drive(*args):
            raise ValueError("path is on a different drive")
        namespace["os"] = types.SimpleNamespace(path=types.SimpleNamespace(relpath=other_drive), replace=os.replace)
    links = []
    package = types.ModuleType("IPython")
    display_module = types.ModuleType("IPython.display")
    display_module.FileLink = lambda path: path
    display_module.display = links.append
    package.display = display_module
    monkeypatch.setitem(sys.modules, "IPython", package)
    monkeypatch.setitem(sys.modules, "IPython.display", display_module)
    exec(cells["download"], namespace)
    assert (tmp_path / "results.zip").is_file()
    assert len(links) == 1
    assert links[0].endswith("results.zip")
