import unicodedata

from openpyxl import load_workbook
import pytest

from src.ocr.application.count_keywords import CountKeywordsUseCase
from src.ocr.infrastructure.postprocessing.keyword_counter import KeywordCounter


def test_exact_vietnamese_counts_keep_accents_and_ignore_case():
    report = KeywordCounter.count_keywords(
        "TRÍ TUỆ NHÂN TẠO, trí tuệ nhân tạo, tri tue nhan tao. Dữ liệu lớn.",
        ["Trí tuệ nhân tạo", "dữ liệu lớn", "học sâu"], "annual", "result.md")
    assert [(item.keyword, item.exact_count, item.count) for item in report.keywords] == [
        ("Trí tuệ nhân tạo", 2, 2), ("dữ liệu lớn", 1, 1), ("học sâu", 0, 0)]
    assert report.total_keywords_searched == 3
    assert report.total_matches == report.total_exact_matches == 3


def test_unicode_blank_and_duplicate_keywords_do_not_inflate_totals(tmp_path):
    path = tmp_path / "keywords.txt"
    path.write_text("\ufeff  Internet di động \n\nINTERNET DI ĐỘNG\n" +
                    unicodedata.normalize("NFD", "Internet di động") + "\n", encoding="utf-8")
    keywords = KeywordCounter.load_keywords(path)
    report = KeywordCounter.count_keywords(unicodedata.normalize("NFD", "Internet di động"), keywords)
    assert keywords == ["Internet di động"]
    assert report.total_keywords_searched == report.total_matches == 1


def test_literal_substrings_independent_phrases_and_non_overlapping_matches():
    report = KeywordCounter.count_keywords("C++ a.b aXb aaaa khai thác dữ liệu lớn", [
        "C++", "a.b", "aa", "dữ liệu", "dữ liệu lớn", "\n", "aa"])
    assert [item.count for item in report.keywords] == [1, 1, 2, 1, 1]
    assert report.total_matches == 6


def test_exact_matching_does_not_join_whitespace_or_fix_ocr():
    report = KeywordCounter.count_keywords("học\nmáy học  máy học may học máy", ["học máy"])
    assert report.total_matches == 1


def test_empty_keyword_list_and_empty_text():
    assert KeywordCounter.count_keywords("body", []).total_matches == 0
    assert KeywordCounter.count_keywords("", ["keyword"]).keywords[0].count == 0


def test_missing_keyword_file_is_not_created(tmp_path):
    missing = tmp_path / "missing.txt"
    with pytest.raises(FileNotFoundError):
        KeywordCounter.load_keywords(missing)
    assert not missing.exists()


def test_empty_keyword_file_is_reported(tmp_path):
    path = tmp_path / "keywords.txt"
    path.write_text("\n \t\n", encoding="utf-8")
    with pytest.raises(ValueError, match="No keywords"):
        KeywordCounter.load_keywords(path)


def test_excel_report_contents_totals_and_refresh_without_changing_markdown(tmp_path):
    document = tmp_path / "annual"
    document.mkdir()
    source = document / "result.md"
    source.write_text("# Trang 1\nĐỔI MỚI, đổi mới. =SUM và C++", encoding="utf-8")
    before = source.read_bytes()
    keywords = tmp_path / "keywords.txt"
    keywords.write_text("Đổi mới\n=SUM\nC++\nKhông có\n", encoding="utf-8")
    report = CountKeywordsUseCase.execute(document, "annual", keywords)
    assert report.total_matches == 4
    workbook = load_workbook(document / "count_keywords.xlsx", rich_text=True)
    try:
        assert workbook.sheetnames == ["Tong_hop", "Trich_dan_chi_tiet"]
        # Sheet 1: Tong_hop (Feature-matrix)
        ws1 = workbook["Tong_hop"]
        assert ws1.cell(1, 1).value == "File_name"
        assert [ws1.cell(1, c).value for c in range(2, 6)] == ["Đổi mới", "=SUM", "C++", "Không có"]
        assert ws1.cell(2, 1).value == "annual"
        assert [ws1.cell(2, c).value for c in range(2, 6)] == [2, 1, 1, 0]

        # Sheet 2: Trich_dan_chi_tiet
        ws2 = workbook["Trich_dan_chi_tiet"]
        assert [ws2.cell(1, c).value for c in range(1, 5)] == [
            "STT", "Tên từ khóa", "Thuộc trang số", "Trích dẫn"]
        # 4 citations: 2 for Đổi mới, 1 for =SUM, 1 for C++
        assert ws2.max_row == 5  # 1 header + 4 rows
        assert [ws2.cell(r, 1).value for r in range(2, 6)] == [1, 2, 3, 4]
        assert [ws2.cell(r, 2).value for r in range(2, 6)] == [
            "Đổi mới", "Đổi mới", "=SUM", "C++"]
        assert [ws2.cell(r, 3).value for r in range(2, 6)] == [1, 1, 1, 1]
    finally:
        workbook.close()
    keywords.write_text("C++\n", encoding="utf-8")
    assert CountKeywordsUseCase.execute(document, "annual", keywords).total_matches == 1
    assert source.read_bytes() == before
    assert not list(document.glob("*.tmp"))


def test_default_keywords_resolve_from_project_when_cwd_changes(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = CountKeywordsUseCase.resolve_keywords_file()
    assert path.is_file()
    assert "Trí tuệ nhân tạo" in KeywordCounter.load_keywords(path)


def test_wheel_uses_bundled_keywords(tmp_path, monkeypatch):
    from src.ocr.application import count_keywords
    path = tmp_path / "keywords.txt"
    path.write_text("học máy", encoding="utf-8")
    monkeypatch.setattr(count_keywords, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(count_keywords, "runtime_resource", lambda name: tmp_path / name)
    assert CountKeywordsUseCase.resolve_keywords_file() == path


def test_missing_result_markdown_is_reported(tmp_path):
    with pytest.raises(FileNotFoundError):
        CountKeywordsUseCase.execute(tmp_path, "missing")
    assert not (tmp_path / "count_keywords.xlsx").exists()


def test_two_sheet_citations_with_pages_and_bold_keyword(tmp_path):
    doc_dir = tmp_path / "doc_test"
    doc_dir.mkdir()
    pages_dir = doc_dir / "pages"
    pages_dir.mkdir()

    # Create page 2 and page 5
    (pages_dir / "page-0002.md").write_text(
        "# Trang 2\nDoanh nghiệp đẩy mạnh chuyển đổi số trong mọi hoạt động sản xuất kinh doanh.",
        encoding="utf-8"
    )
    (pages_dir / "page-0005.md").write_text(
        "# Trang 5\nỨng dụng trí tuệ nhân tạo và chuyển đổi số mang lại hiệu quả vượt bậc.",
        encoding="utf-8"
    )
    # result.md combines both
    (doc_dir / "result.md").write_text(
        "# Trang 2\nDoanh nghiệp đẩy mạnh chuyển đổi số trong mọi hoạt động sản xuất kinh doanh.\n---\n"
        "# Trang 5\nỨng dụng trí tuệ nhân tạo và chuyển đổi số mang lại hiệu quả vượt bậc.",
        encoding="utf-8"
    )

    kw_file = tmp_path / "keywords.txt"
    kw_file.write_text("Trí tuệ nhân tạo\nChuyển đổi số\nĐiện toán đám mây\n", encoding="utf-8")

    report = CountKeywordsUseCase.execute(doc_dir, "doc_test", kw_file)
    assert report.total_matches == 3  # 1 Trí tuệ nhân tạo + 2 Chuyển đổi số + 0 Điện toán đám mây
    assert len(report.citations) == 3

    wb = load_workbook(doc_dir / "count_keywords.xlsx", rich_text=True)
    try:
        assert wb.sheetnames == ["Tong_hop", "Trich_dan_chi_tiet"]
        ws1 = wb["Tong_hop"]
        assert ws1.cell(1, 1).value == "File_name"
        assert ws1.cell(1, 2).value == "Trí tuệ nhân tạo"
        assert ws1.cell(1, 3).value == "Chuyển đổi số"
        assert ws1.cell(1, 4).value == "Điện toán đám mây"
        assert ws1.cell(2, 1).value == "doc_test"
        assert ws1.cell(2, 2).value == 1
        assert ws1.cell(2, 3).value == 2
        assert ws1.cell(2, 4).value == 0  # 0 count column is retained!

        ws2 = wb["Trich_dan_chi_tiet"]
        # Order must be: Trí tuệ nhân tạo (p5), Chuyển đổi số (p2), Chuyển đổi số (p5)
        # Because "Trí tuệ nhân tạo" is first in keywords.txt, then "Chuyển đổi số"
        # and for Chuyển đổi số, page 2 < page 5
        rows = [
            (ws2.cell(r, 1).value, ws2.cell(r, 2).value, ws2.cell(r, 3).value)
            for r in range(2, 5)
        ]
        assert rows == [
            (1, "Trí tuệ nhân tạo", 5),
            (2, "Chuyển đổi số", 2),
            (3, "Chuyển đổi số", 5),
        ]

        # Verify rich text citation formatting for row 2 (Trí tuệ nhân tạo)
        cit_cell = ws2.cell(2, 4).value
        # Should be a CellRichText containing bold TextBlock
        bold_blocks = [block for block in cit_cell if getattr(getattr(block, "font", None), "b", False)]
        assert len(bold_blocks) >= 1
        assert bold_blocks[0].text.casefold() == "trí tuệ nhân tạo"
    finally:
        wb.close()

