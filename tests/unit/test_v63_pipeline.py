"""V6.3 fault injection/contract tests; simulated receipts are not GPU evidence."""
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
from unittest.mock import Mock

import openpyxl
import pymupdf
import pytest

from src.ocr.application.process_batch import process_batch, input_snapshot
from src.ocr.application.process_document import process_document, validate_gpu_source
from src.ocr.domain.models.processing import PublicationError
from src.ocr.infrastructure.export.completion import verify_completed
from src.ocr.infrastructure.postprocessing.keyword_contract import dictionary_snapshot, count_report, validate_report
from src.ocr.infrastructure.export.keyword_v63 import matrix_bytes, verify_matrix
from src.ocr.infrastructure.pdf.ocr_bridge import ocr_bridge
from src.ocr.runtime.atomic_io import atomic_write_json, atomic_write_bytes
from src.ocr.runtime.gpu_scheduler import sha256, GPULease
from src.ocr.runtime.paths import PROJECT_ROOT
from test_document_furniture_v3 import make_doc, block

KEYWORDS = PROJECT_ROOT/"config/keywords.txt"


def pdf(path):
    with pymupdf.open() as doc:
        doc.new_page().insert_text((40,100),"NATIVE DECOY MUST NEVER ENTER OCR COUNTS")
        doc.save(path)


def simulated_backend(request, **kwargs):
    """Only a finalizer test double. Real acceptance must use actual GPU phases."""
    from src.ocr.application.export_document import ExportDocumentUseCase
    doc = make_doc([[block("Chuyển đổi số Internet di động Internet di động",1,200)]])
    doc.document_id = Path(request["pdf"]).stem
    doc.file_path, doc.file_hash = request["pdf"],request["pdf_sha256"][:16]
    doc.quality.document_id = doc.document_id
    doc.manifest.update(document_id=doc.document_id,file_hash=doc.file_hash,total_pages=1,
        configuration={"ocr_source_policy":"gpu_ocr_only_v1"},
        gpu_runtime={"phase_a_placement":{"detector":{"expected_provider":"CUDAExecutionProvider"},
                                         "recognizer":{"device":"cuda","workload":"document"}}})
    doc.pages[0].metadata.update(gpu_ocr={"scope":"all_pages","source_policy":"gpu_ocr_only_v1",
        "detector_device":"cuda","recognizer_device":"cuda","detector_calls":1,"recognizer_forward_calls":1,
        "workload":"document","status":"OCR_PROCESSED"},
        ocr_structure_bridge={"policy":"raster-plus-gpu-text-v1"}, test_double_not_gpu_acceptance=True)
    for item in doc.pages[0].evidence_blocks:
        item.source = "ocr"
    ExportDocumentUseCase.execute(doc,Path(request["output"]),assets_mode="none",count_keywords=False)


def processor_factory(calls):
    def processor(record,snapshot,output,**options):
        def backend(request,**kwargs):
            calls.append(record["id"])
            simulated_backend(request,**kwargs)
        request = {"pdf":record["path"],"pdf_sha256":record["pdf_sha256"],"output":str(output),"ocr_identity":record["pdf_sha256"]}
        return process_document(request,snapshot,resume=options["resume"],backend=backend)
    return processor


def setup_inputs(tmp_path, count=3):
    source = tmp_path/"inputs"
    source.mkdir()
    pdf(source/"a.PDF")
    for i in range(1,count):
        (source/f"{chr(97+i)}.pdf").write_bytes((source/"a.PDF").read_bytes())
    return source,tmp_path/"output"


def test_canonical_display_and_validation(tmp_path):
    snapshot = dictionary_snapshot(KEYWORDS)
    assert len(snapshot["canonical_keywords"]) == 87 and len(snapshot["display_columns"]) == 88
    path = tmp_path/"result.md"
    path.write_text("Internet di động Internet di động chuyển đổi số",encoding="utf-8")
    report = count_report("test","abc",path,snapshot,[1])
    duplicate = [i for i,k in enumerate(snapshot["display_columns"]) if k.casefold()=="internet di động"]
    assert len(duplicate)==2 and [report["display_counts"][i] for i in duplicate]==[2,2]
    assert report["total_display_matches"]==report["total_canonical_matches"]+2
    validate_report(report,snapshot,path)
    report["canonical_counts"].pop(next(iter(report["canonical_counts"])))
    with pytest.raises(ValueError,match="canonical"):
        validate_report(report,snapshot,path)


def test_three_pdf_order_resume_matrix_repair_and_frozen_inputs(tmp_path):
    source,output=setup_inputs(tmp_path)
    calls=[]
    result=process_batch(source,output,{},KEYWORDS,processor_factory(calls))
    assert calls==["a","b","c"]
    assert result["published_revision"]==3
    with sqlite3.connect(output/"batch_state.sqlite") as store:
        events=list(store.execute("SELECT document_id,stage FROM events ORDER BY sequence"))
    assert events.index(("a","DOCUMENT_COMMITTED")) < events.index(("b","OCR_RUNNING"))
    assert events.index(("b","DOCUMENT_COMMITTED")) < events.index(("c","OCR_RUNNING"))
    assert all(verify_completed(output/name) for name in calls)
    matrix=Path(result["matrix"])
    matrix.write_bytes(b"damaged Excel")
    pdf(source/"added-after-snapshot.pdf")
    resumed=process_batch(source,output,{},KEYWORDS,processor_factory(calls),resume=True)
    assert calls==["a","b","c"] and resumed["published_revision"]==3
    assert len(resumed["documents"])==3
    snapshot=dictionary_snapshot(KEYWORDS)
    verify_matrix(matrix,resumed["batch_id"],3,resumed["documents"],snapshot)


def test_excel_replace_failure_stops_then_resumes_without_ocr(tmp_path,monkeypatch):
    source,output=setup_inputs(tmp_path)
    from src.ocr.application import publish_keyword_matrix as publisher
    original=publisher.atomic_write_bytes
    writes=[]
    def fail_second(path,data):
        writes.append(path)
        if len(writes)==2:
            raise PermissionError("Excel open")
        original(path,data)
    monkeypatch.setattr(publisher,"atomic_write_bytes",fail_second)
    calls=[]
    with pytest.raises(PermissionError):
        process_batch(source,output,{},KEYWORDS,processor_factory(calls))
    assert calls==["a"] and verify_completed(output/"a")
    with sqlite3.connect(output/"batch_state.sqlite") as store:
        meta=dict(store.execute("SELECT key,value FROM meta"))
        assert json.loads(meta["prepared_revision"])==1 and json.loads(meta["published_revision"])==0
    monkeypatch.setattr(publisher,"atomic_write_bytes",original)
    process_batch(source,output,{},KEYWORDS,processor_factory(calls),resume=True)
    assert calls==["a","b","c"]


def test_per_document_excel_failure_resume_keeps_counter_and_ocr_sealed(tmp_path,monkeypatch):
    source,output=setup_inputs(tmp_path,1)
    import src.ocr.application.process_document as module
    original=module.export_document
    monkeypatch.setattr(module,"export_document",Mock(side_effect=PermissionError("busy workbook")))
    calls=[]
    with pytest.raises(PublicationError):
        process_batch(source,output,{},KEYWORDS,processor_factory(calls))
    assert calls==["a"] and (output/"a/keyword_counts.json").is_file()
    monkeypatch.setattr(module,"export_document",original)
    monkeypatch.setattr(module,"count_report",Mock(side_effect=AssertionError("must reuse sealed report")))
    process_batch(source,output,{},KEYWORDS,processor_factory(calls),resume=True)
    assert calls==["a"]


def test_failure_blank_row_and_retry_replaces_same_row(tmp_path):
    source,output=setup_inputs(tmp_path,2)
    calls=[]
    good=processor_factory(calls)
    def fail(record,*args,**kwargs):
        if record["id"]=="a":
            raise ValueError("invalid document fixture")
        return good(record,*args,**kwargs)
    result=process_batch(source,output,{},KEYWORDS,fail)
    assert result["documents"][0]["counts"] is None
    process_batch(source,output,{},KEYWORDS,good,resume=True)
    assert calls==["b"]
    result=process_batch(source,output,{},KEYWORDS,good,resume=True,retry_failed=True)
    assert calls==["b","a"] and len(result["documents"])==2


def test_native_text_never_copied_into_structure_bridge(tmp_path):
    path=tmp_path/"source.pdf"
    pdf(path)
    extracted=make_doc([[block("ONLY OCR WORDS",1,200)]]).pages[0]
    with pymupdf.open(path) as source:
        with ocr_bridge(source[0],extracted,120) as bridge:
            text=" ".join(bridge[0].get_text().split())
            assert "ONLY OCR WORDS" in text and "NATIVE DECOY" not in text
    extracted.blocks[0].source="native"
    with pymupdf.open(path) as source, pytest.raises(ValueError,match="native"):
        with ocr_bridge(source[0],extracted,120):
            pass


def test_500_row_projection_has_no_50_document_limit(tmp_path):
    snapshot=dictionary_snapshot(KEYWORDS)
    records=[{"id":str(i),"name":f"{i}.pdf","pdf_sha256":str(i),"counts":[i]*88,"outcome":"SUCCEEDED"} for i in range(500)]
    path=tmp_path/"matrix.xlsx"
    atomic_write_bytes(path,matrix_bytes("test",500,records,snapshot))
    verify_matrix(path,"test",500,records,snapshot)
    book=openpyxl.load_workbook(path,read_only=True)
    try:
        assert len(list(book["Ma trận từ khóa"].iter_rows()))==502
    finally:
        book.close()


def test_cpu_parent_and_cli_rejections():
    code="import sys; import src.ocr.cli.unified,src.ocr.application.process_batch; assert 'torch' not in sys.modules"
    subprocess.run([sys.executable,"-c",code],cwd=PROJECT_ROOT,check=True)
    result=subprocess.run([sys.executable,"run_batch.py","--backend","gpu_two_phase","--device","cpu"],cwd=PROJECT_ROOT,capture_output=True)
    assert result.returncode!=0 and b"invalid choice" in result.stderr


def test_replace_then_ack_crash_reconciles_without_rewriting_same_revision(tmp_path,monkeypatch):
    source,output=setup_inputs(tmp_path,2)
    from src.ocr.infrastructure.state.batch_store import BatchStore
    from src.ocr.application import publish_keyword_matrix as publisher
    acknowledge=BatchStore.acknowledge
    def crash(store,revision,checksum):
        if revision==1:
            raise RuntimeError("crash after replace before acknowledgement")
        return acknowledge(store,revision,checksum)
    monkeypatch.setattr(BatchStore,"acknowledge",crash)
    calls=[]
    with pytest.raises(RuntimeError,match="crash after replace"):
        process_batch(source,output,{},KEYWORDS,processor_factory(calls))
    assert calls==["a"]
    monkeypatch.setattr(BatchStore,"acknowledge",acknowledge)
    writes=Mock(wraps=publisher.atomic_write_bytes)
    monkeypatch.setattr(publisher,"atomic_write_bytes",writes)
    process_batch(source,output,{},KEYWORDS,processor_factory(calls),resume=True)
    assert calls==["a","b"] and writes.call_count==1  # Only document B needs a new workbook.


def test_missing_entire_keyword_group_recovers_without_ocr(tmp_path):
    from src.ocr.domain.models.processing import KEYWORD_ARTIFACTS
    from src.ocr.infrastructure.export.artifact_validator import validate_document_artifacts
    source,output=setup_inputs(tmp_path,1)
    calls=[]
    processor=processor_factory(calls)
    process_batch(source,output,{},KEYWORDS,processor)
    directory=output/"a"
    for name in KEYWORD_ARTIFACTS:
        (directory/name).unlink()
    manifest=json.loads((directory/"manifest.json").read_text(encoding="utf-8"))
    manifest.pop("keyword_contract")
    atomic_write_json(directory/"manifest.json",manifest)
    assert validate_document_artifacts(directory)["status"]=="FAIL"
    result=process_batch(source,output,{},KEYWORDS,processor,resume=True)
    assert calls==["a"] and result["published_revision"]==2
    assert verify_completed(directory)


def test_recount_uses_same_counter_and_matrix_without_ocr(tmp_path):
    source,output=setup_inputs(tmp_path,2)
    calls=[]
    processor=processor_factory(calls)
    first=process_batch(source,output,{},KEYWORDS,processor)
    result=process_batch(source,output,{},KEYWORDS,processor,resume=True,force_recount=True)
    assert calls==["a","b"] and result["published_revision"]==4
    assert [r["counts"] for r in first["documents"]]==[r["counts"] for r in result["documents"]]


def test_changed_pdf_and_dictionary_refuse_snapshot_reuse(tmp_path):
    from src.ocr.domain.models.processing import IdentityMismatch
    source,output=setup_inputs(tmp_path,1)
    calls=[]
    process_batch(source,output,{},KEYWORDS,processor_factory(calls))
    original=(source/"a.PDF").read_bytes()
    (source/"a.PDF").write_bytes(original+b"\n% changed bytes")
    with pytest.raises(IdentityMismatch,match="INPUT_CHANGED"):
        process_batch(source,output,{},KEYWORDS,processor_factory(calls),resume=True)
    (source/"a.PDF").write_bytes(original)
    keywords=tmp_path/"keywords.txt"
    keywords.write_text(KEYWORDS.read_text(encoding="utf-8").upper(),encoding="utf-8")
    with pytest.raises(ValueError,match="IDENTITY_MISMATCH"):
        process_batch(source,output,{},keywords,processor_factory(calls),resume=True)
    assert calls==["a"]


def test_batch_exclusive_kernel_lock(tmp_path):
    source,output=setup_inputs(tmp_path,1)
    with GPULease(output/"batch.lock",timeout=0):
        with pytest.raises(TimeoutError,match="LEASE_BUSY"):
            process_batch(source,output,{},KEYWORDS,processor_factory([]))
    assert not (output/"batch_state.sqlite").exists()


def test_source_validator_rejects_native_or_synthetic_probe(tmp_path):
    source,output=setup_inputs(tmp_path,1)
    calls=[]
    process_batch(source,output,{},KEYWORDS,processor_factory(calls))
    page_path=output/"a/pages/page-0001.json"
    page=json.loads(page_path.read_text(encoding="utf-8"))
    page["blocks"][0]["source"]="native"
    atomic_write_json(page_path,page)
    with pytest.raises(ValueError,match="Native text"):
        validate_gpu_source(output/"a")
    page["blocks"][0]["source"]="primary_ocr"
    atomic_write_json(page_path,page)
    manifest_path=output/"a/manifest.json"
    manifest=json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["gpu_runtime"]["phase_a_placement"]["recognizer"]["workload"]="synthetic_backend_probe_native_only_document"
    atomic_write_json(manifest_path,manifest)
    with pytest.raises(ValueError,match="real document"):
        validate_gpu_source(output/"a")


def test_500_document_coordinator_has_no_50_item_cutoff(tmp_path,monkeypatch):
    """Logic scale only: no OCR/GPU and no 500-revision disk benchmark claim."""
    import src.ocr.application.process_batch as coordinator
    template=tmp_path/"template.pdf"
    pdf(template)
    source=tmp_path/"inputs"
    source.mkdir()
    data=template.read_bytes()
    for i in range(500):
        (source/f"doc{i:04d}.pdf").write_bytes(data)
    class MemoryStore:
        current=None
        def __init__(self,path,identity,records):
            self.meta={"identity":identity,"prepared_revision":0,"published_revision":-1}
            self.rows={r["id"]:r for r in records}
            MemoryStore.current=self
        def get(self,key): return self.meta[key]
        def records(self): return list(self.rows.values())
        def event(self,*args,**kwargs): pass
        def prepare(self,record,**kwargs):
            assert self.get("prepared_revision")==self.get("published_revision")
            self.rows[record["id"]]=record
            self.meta["prepared_revision"]+=1
        def close(self): pass
        def export_events(self,path): pass
    def publication(store,*args):
        store.meta["published_revision"]=store.get("prepared_revision")
    monkeypatch.setattr(coordinator,"BatchStore",MemoryStore)
    monkeypatch.setattr(coordinator,"publish_keyword_matrix",publication)
    monkeypatch.setattr("src.ocr.infrastructure.export.completion.verify_completed",lambda path:True)
    calls=[]
    def processor(record,snapshot,output,**kwargs):
        assert MemoryStore.current.get("published_revision")==len(calls)
        calls.append(record["id"])
        directory=output/record["id"]
        directory.mkdir()
        (directory/"result.md").write_text("",encoding="utf-8")
        atomic_write_json(directory/"keyword_counts.json",count_report(record["id"],record["pdf_sha256"],directory/"result.md",snapshot,[1]))
        atomic_write_json(directory/"COMPLETED.json",{"unit_fixture":True})
        return {"outcome":"SUCCEEDED","output":str(directory),"report_sha256":sha256(directory/"keyword_counts.json"),
                "completion_sha256":sha256(directory/"COMPLETED.json"),"validation":"NEEDS_REVIEW"}
    result=process_batch(source,tmp_path/"output",{},KEYWORDS,processor)
    assert len(calls)==500 and len(result["documents"])==500 and result["published_revision"]==500


def test_legacy_count_cannot_bypass_v63_publication_gate(tmp_path):
    from src.ocr.application.count_keywords import CountKeywordsUseCase
    source,output=setup_inputs(tmp_path,1)
    process_batch(source,output,{},KEYWORDS,processor_factory([]))
    before=sha256(output/"a/count_keywords.xlsx")
    with pytest.raises(ValueError,match="publication gate"):
        CountKeywordsUseCase.execute(output/"a","a")
    assert sha256(output/"a/count_keywords.xlsx")==before
    assert verify_completed(output/"a")
    events=[json.loads(line) for line in (output/"logs/batch.jsonl").read_text(encoding="utf-8").splitlines()]
    assert events[-1]["stage"]=="DOCUMENT_COMMITTED"


def test_failed_force_ocr_clears_old_counts_and_retry_clears_error(tmp_path):
    source,output=setup_inputs(tmp_path,1)
    processor=processor_factory([])
    first=process_batch(source,output,{},KEYWORDS,processor)
    assert first["documents"][0]["total_canonical_matches"]>0
    def broken(*args,**kwargs):
        raise ValueError("forced OCR failed")
    failed=process_batch(source,output,{},KEYWORDS,broken,resume=True,force_ocr=True)
    record=failed["documents"][0]
    assert record["counts"] is None and record["total_canonical_matches"]==0
    assert record["validation"] is None and "report_sha256" not in record
    snapshot=dictionary_snapshot(KEYWORDS)
    verify_matrix(failed["matrix"],failed["batch_id"],failed["published_revision"],failed["documents"],snapshot)
    recovered=process_batch(source,output,{},KEYWORDS,processor,resume=True,retry_failed=True)
    assert recovered["documents"][0]["outcome"]=="SUCCEEDED"
    assert "error" not in recovered["documents"][0]


def test_failed_fresh_extraction_invalidates_previous_ocr_receipt(tmp_path):
    from src.ocr.application.process_document import valid_ocr
    source,output=setup_inputs(tmp_path,1)
    first=process_batch(source,output,{},KEYWORDS,processor_factory([]))
    item=first["documents"][0]
    request={"pdf":item["path"],"pdf_sha256":item["pdf_sha256"],"output":str(output),"ocr_identity":item["pdf_sha256"]}
    assert valid_ocr(output/"a",item["pdf_sha256"])
    def broken(request,**kwargs):
        assert kwargs["resume"] is False
        assert not (output/"a/ocr_receipt.json").exists()
        assert not verify_completed(output/"a")
        raise ValueError("new inference failed")
    with pytest.raises(ValueError,match="new inference failed"):
        process_document(request,dictionary_snapshot(KEYWORDS),resume=False,backend=broken)
    assert not valid_ocr(output/"a",item["pdf_sha256"])


def test_duplicate_sources_keep_rows_and_visible_duplicate_flag(tmp_path):
    source,output=setup_inputs(tmp_path,1)
    (source/"b.pdf").write_bytes((source/"a.PDF").read_bytes())
    result=process_batch(source,output,{},KEYWORDS,processor_factory([]))
    assert len(result["documents"])==2 and result["documents"][1]["duplicate_of"]=="a"
    book=openpyxl.load_workbook(result["matrix"],read_only=True,data_only=True)
    try:
        rows=list(book["Trạng thái"].iter_rows(values_only=True))
        assert rows[0][-1]=="Duplicate of" and rows[2][-1]=="a"
        assert dict(book["Tổng quan"].iter_rows(values_only=True))["Nguồn hash duy nhất"]==1
    finally:
        book.close()


@pytest.mark.parametrize("mutation",["pdf","runtime"])
def test_mid_document_identity_change_blocks_publication(tmp_path,mutation):
    from src.ocr.domain.models.processing import IdentityMismatch
    source,output=setup_inputs(tmp_path,2)
    calls=[]
    inner=processor_factory(calls)
    changed=False
    def check():
        if changed:
            raise IdentityMismatch("Runtime changed during OCR")
    def processor(record,*args,**kwargs):
        nonlocal changed
        result=inner(record,*args,**kwargs)
        if mutation=="pdf":
            path=Path(record["path"])
            path.write_bytes(path.read_bytes()+b'\n')
        else:
            changed=True
        return result
    with pytest.raises(IdentityMismatch):
        process_batch(source,output,{},KEYWORDS,processor,identity_check=check)
    assert calls==["a"]
    with sqlite3.connect(output/"batch_state.sqlite") as store:
        assert json.loads(store.execute("select value from meta where key='published_revision'").fetchone()[0])==0
        assert store.execute("select count(*) from events where stage='DOCUMENT_COMMITTED'").fetchone()[0]==0
