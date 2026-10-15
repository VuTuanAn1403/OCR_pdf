import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pymupdf
import pytest

from src.ocr.application.convert_document import ConvertDocumentUseCase
from src.ocr.application.count_keywords import CountKeywordsUseCase
from src.ocr.cli.batch_runner import BatchOCRRunner
from src.ocr.infrastructure.export.completion import verify_completed
from src.ocr.infrastructure.export.result_reader import load_document_json

ROOT = Path(__file__).resolve().parents[2]


def _native_pdf(path):
    with pymupdf.open() as source:
        page = source.new_page()
        page.insert_textbox((40, 40, 550, 700), (
            "Annual financial report for shareholders and investors. "
            "The company uses artificial intelligence and cloud computing. "
            "We invest in artificial intelligence for business growth. "
            "Management develops technology to improve customer service and operations. "
            "Revenue and profit are reviewed by the board each year.\n") * 4)
        source.save(path)


def _excel_counts(directory):
    workbook = load_workbook(directory / "count_keywords.xlsx", data_only=True)
    try:
        sheet = workbook.active
        return {sheet.cell(r, 2).value: sheet.cell(r, 3).value
                for r in range(5, sheet.max_row)}
    finally:
        workbook.close()


@pytest.mark.parametrize("mode", ["legacy", "compact"])
def test_native_pipeline_exports_keyword_excel_and_preserves_completion(tmp_path, mode):
    pdf = tmp_path / "annual.pdf"
    _native_pdf(pdf)
    keywords = tmp_path / "keywords.txt"
    keywords.write_text("artificial intelligence\ncloud computing\nabsent\n", encoding="utf-8")
    with ConvertDocumentUseCase() as converter:
        document = converter.execute(str(pdf), profile_name="fast", output_dir=str(tmp_path / "out"),
            assets_mode="none", output_mode=mode, keywords_file=str(keywords),
            cache_dir=str(tmp_path / "cache"))
    directory = tmp_path / "out/annual"
    assert len(document.pages) == 1
    assert document.quality.pages_processed_native == 1
    assert _excel_counts(directory) == {"artificial intelligence": 8, "cloud computing": 4, "absent": 0}
    assert len(load_document_json(directory)["pages"]) == 1
    assert verify_completed(directory)
    keywords.write_text("cloud computing\n", encoding="utf-8")
    CountKeywordsUseCase.execute(directory, "annual", keywords)
    assert _excel_counts(directory) == {"cloud computing": 4}
    assert verify_completed(directory)
    (directory / "result.md").write_text("changed without re-export", encoding="utf-8")
    CountKeywordsUseCase.execute(directory, "annual", keywords)
    assert not verify_completed(directory)


@pytest.mark.parametrize("workers", [1, 2])
def test_batch_counts_custom_keywords_and_refreshes_cached_outputs(tmp_path, workers):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    for name in ("annual", "second"):
        _native_pdf(inputs / f"{name}.pdf")
    keywords = tmp_path / "keywords.txt"
    keywords.write_text("artificial intelligence\n", encoding="utf-8")
    options = dict(input_dir=str(inputs), output_dir=str(tmp_path / "out"), profile="fast",
                   assets_mode="none", output_mode="compact", workers=workers,
                   keywords_file=str(keywords), cache_dir=str(tmp_path / "cache"), excel_path=None)
    BatchOCRRunner(**options).execute_batch()
    directory = tmp_path / "out/annual"
    assert _excel_counts(directory) == {"artificial intelligence": 8}
    original_result = (directory / "result.md").read_bytes()
    keywords.write_text("cloud computing\n", encoding="utf-8")
    BatchOCRRunner(**options).execute_batch()
    manifest = json.loads((tmp_path / "out/batch_manifest.json").read_text(encoding="utf-8"))
    assert all(entry["status"] == "CACHED" for entry in manifest)
    for name in ("annual", "second"):
        assert _excel_counts(tmp_path / "out" / name) == {"cloud computing": 4}
    assert (directory / "result.md").read_bytes() == original_result
    assert verify_completed(directory)


def test_count_only_command_runs_from_another_directory(tmp_path):
    directory = tmp_path / "annual"
    directory.mkdir()
    (directory / "result.md").write_text("ĐỔI MỚI, đổi mới.", encoding="utf-8")
    keywords = tmp_path / "keywords.txt"
    keywords.write_text("Đổi mới", encoding="utf-8")
    result = subprocess.run([sys.executable, str(ROOT / "scripts/run_count_only.py"), str(directory),
                             "--keywords-file", str(keywords)], cwd=tmp_path,
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stderr
    assert _excel_counts(directory) == {"Đổi mới": 2}
