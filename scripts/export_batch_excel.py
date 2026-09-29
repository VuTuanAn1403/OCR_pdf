#!/usr/bin/env python3
import argparse
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)


def load_json(path: Path, default: Any = None) -> Any:
    try:
        with path.open("r", encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, json.JSONDecodeError):
        return default


def style_table(ws, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def fit_columns(ws, min_width: int = 10, max_width: int = 55) -> None:
    for idx, column in enumerate(ws.columns, 1):
        values = [str(cell.value) if cell.value is not None else "" for cell in column[:200]]
        width = max([len(value) for value in values] + [min_width]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(max_width, width)


def append_rows(ws, headers: List[str], rows: Iterable[List[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)


def build_workbook(output_dir: Path, excel_path: Path) -> None:
    batch = load_json(output_dir / "batch_report.json", {}) or {}
    summary = batch.get("summary", {})
    files = batch.get("files", [])
    review_queue = load_json(output_dir / "review_queue.json", []) or []

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tong quan"

    overview["A1"] = f"BÁO CÁO KIỂM NGHIỆM OCR BATCH - {len(files)} PDF"
    overview["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    overview["A1"].fill = HEADER_FILL
    overview.merge_cells("A1:D1")
    overview["A2"] = "Pipeline"
    overview["B2"] = summary.get("pipeline_version")
    overview["C2"] = "Profile"
    overview["D2"] = summary.get("profile")

    metrics = [
        ("Tổng file", summary.get("total_files")),
        ("File hoàn tất", summary.get("completed")),
        ("File lỗi", summary.get("failed")),
        ("Tổng trang", summary.get("processed_pages")),
        ("Tổng thời gian (giây)", summary.get("total_batch_duration_seconds")),
        ("Trung bình giây/trang", summary.get("average_seconds_per_page")),
        ("Trung vị giây/trang", summary.get("median_seconds_per_page")),
        ("P95 giây/trang", summary.get("p95_seconds_per_page")),
        ("Trang/phút", summary.get("overall_pages_per_minute")),
        ("File nhanh nhất", (summary.get("fastest_file") or {}).get("file_name")),
        ("File chậm nhất", (summary.get("slowest_file") or {}).get("file_name")),
    ]
    overview.append([])
    overview.append(["Chỉ số", "Giá trị"])
    for label, value in metrics:
        overview.append([label, value])
    for cell in overview[4]:
        cell.fill = SUBHEADER_FILL
        cell.font = Font(bold=True)

    overview["F2"] = "File"
    overview["G2"] = "Giây/trang"
    completed_files = [f for f in files if f.get("status") in ("COMPLETED", "CACHED") and f.get("seconds_per_page") is not None]
    for row_idx, item in enumerate(completed_files, 3):
        overview.cell(row=row_idx, column=6, value=item.get("document_id"))
        overview.cell(row=row_idx, column=7, value=item.get("seconds_per_page"))
    if completed_files:
        chart = BarChart()
        chart.title = "Tốc độ theo tài liệu"
        chart.y_axis.title = "Giây/trang"
        chart.x_axis.title = "Tài liệu"
        chart.height = 14
        chart.width = 24
        chart.add_data(Reference(overview, min_col=7, min_row=2, max_row=2 + len(completed_files)), titles_from_data=True)
        chart.set_categories(Reference(overview, min_col=6, min_row=3, max_row=2 + len(completed_files)))
        overview.add_chart(chart, "I2")
    overview.column_dimensions["A"].width = 30
    overview.column_dimensions["B"].width = 28
    overview.column_dimensions["C"].width = 18
    overview.column_dimensions["D"].width = 18
    overview.column_dimensions["F"].width = 32
    overview.column_dimensions["G"].width = 14

    file_sheet = workbook.create_sheet("Chi tiet file")
    file_headers = [
        "File", "Dung lượng MB", "Trang", "Trạng thái", "Thời gian (s)",
        "Giây/trang", "Trang/phút", "Native", "OCR", "Mixed", "Fallback lines",
        "Tổng lines", "Fallback %", "Confidence", "Quality", "Peak RAM MB",
        "Phân loại", "Unicode lỗi", "Xoay ảnh", "Extraction lỗi", "Số bảng",
    ]
    file_rows = []
    page_rows: List[List[Any]] = []
    review_rows: List[List[Any]] = []

    for item in files:
        document_id = item.get("document_id")
        quality = load_json(output_dir / document_id / "quality.json", {}) or {}
        result = load_json(output_dir / document_id / "result.json", {}) or {}
        result_pages = {p.get("page_num"): p for p in result.get("pages", [])}
        reports = quality.get("page_reports", [])
        unicode_pages = sum(bool(p.get("unicode_corruption_detected")) for p in reports)
        rotated_pages = sum(bool(p.get("orientation_fallback")) for p in reports)
        extraction_failures = sum(bool(p.get("native_extraction_failed")) for p in reports)
        table_count = sum(len(p.get("tables", [])) for p in result.get("pages", []))

        file_rows.append([
            item.get("file_name"), item.get("file_size_mb"), item.get("total_pages"),
            item.get("status"), item.get("wall_clock_seconds"), item.get("seconds_per_page"),
            item.get("pages_per_minute"), item.get("pages_native"), item.get("pages_ocr"),
            item.get("pages_mixed"), item.get("fallback_lines"), item.get("total_lines"),
            item.get("fallback_ratio"), item.get("mean_confidence"), item.get("mean_quality_score"),
            item.get("peak_memory_mb"), ", ".join(item.get("classifications", [])), unicode_pages,
            rotated_pages, extraction_failures, table_count,
        ])

        for page in reports:
            page_num = page.get("page")
            page_result = result_pages.get(page_num, {})
            suspicious = page.get("suspicious_lines", []) or []
            page_rows.append([
                item.get("file_name"), page_num, page.get("page_type"),
                page.get("time_taken_seconds"), page.get("native_text_score"),
                page.get("ocr_confidence"), page.get("quality_score"), page.get("total_blocks"),
                page.get("fallback_count"), len(suspicious), bool(page.get("unicode_corruption_detected")),
                bool(page.get("native_extraction_failed")), page.get("orientation_fallback"),
                bool(page.get("vector_table_probe_enabled")), len(page_result.get("tables", [])),
                " | ".join(suspicious),
            ])
            reasons = []
            if page.get("page_type") != "PRODUCT_SHOWCASE_COLLAGE":
                if page.get("quality_score", 1.0) < 0.90:
                    reasons.append("Quality < 0.90")
                if page.get("ocr_confidence", 1.0) < 0.90:
                    reasons.append("Confidence < 0.90")
            if page.get("unicode_corruption_detected"):
                reasons.append("Unicode lỗi")
            if page.get("native_extraction_failed"):
                reasons.append("Không đọc được native text")
            if page.get("time_taken_seconds", 0.0) > 25.0:
                reasons.append("Chậm > 25 giây")
            if page.get("total_blocks", 0) == 0 and len(page_result.get("tables", [])) == 0:
                reasons.append("Không có block nội dung")
            if reasons:
                review_rows.append([
                    item.get("file_name"), page_num, ", ".join(reasons), page.get("page_type"),
                    page.get("time_taken_seconds"), page.get("ocr_confidence"),
                    page.get("quality_score"), len(suspicious), " | ".join(suspicious),
                ])

    append_rows(file_sheet, file_headers, file_rows)
    style_table(file_sheet)
    fit_columns(file_sheet)
    for row in range(2, file_sheet.max_row + 1):
        for col in (13, 14, 15):
            c = file_sheet.cell(row, col)
            if c.value is not None:
                c.number_format = "0.00%"
        status_cell = file_sheet.cell(row, 4)
        if status_cell.value in ("COMPLETED", "CACHED"):
            status_cell.fill = GOOD_FILL
        elif status_cell.value == "PENDING":
            status_cell.fill = WARN_FILL
        else:
            status_cell.fill = BAD_FILL

    page_sheet = workbook.create_sheet("Chi tiet trang")
    page_headers = [
        "File", "Trang", "Route", "Thời gian (s)", "Native score", "OCR confidence",
        "Quality", "Blocks", "Fallback", "Dòng nghi vấn", "Unicode lỗi",
        "Extraction lỗi", "Xoay ảnh", "Vector table probe", "Số bảng", "Chi tiết nghi vấn",
    ]
    append_rows(page_sheet, page_headers, page_rows)
    style_table(page_sheet)
    fit_columns(page_sheet, max_width=70)
    for row in range(2, page_sheet.max_row + 1):
        for col in (5, 6, 7):
            page_sheet.cell(row, col).number_format = "0.00%"
    if page_sheet.max_row > 1:
        page_sheet.conditional_formatting.add(
            f"G2:G{page_sheet.max_row}",
            ColorScaleRule(start_type="num", start_value=0.8, start_color="F8696B",
                           mid_type="num", mid_value=0.95, mid_color="FFEB84",
                           end_type="num", end_value=1.0, end_color="63BE7B"),
        )

    review_sheet = workbook.create_sheet("Can kiem tra")
    review_headers = [
        "File", "Trang", "Lý do", "Route", "Thời gian (s)", "Confidence",
        "Quality", "Dòng nghi vấn", "Chi tiết",
    ]
    append_rows(review_sheet, review_headers, review_rows)
    style_table(review_sheet)
    fit_columns(review_sheet, max_width=75)
    for row in range(2, review_sheet.max_row + 1):
        review_sheet.cell(row, 3).fill = WARN_FILL
        review_sheet.cell(row, 6).number_format = "0.00%"
        review_sheet.cell(row, 7).number_format = "0.00%"

    queue_sheet = workbook.create_sheet("Review queue")
    append_rows(
        queue_sheet,
        ["Mức độ", "File", "Lý do"],
        [[entry.get("priority"), entry.get("file"), entry.get("reason")] for entry in review_queue],
    )
    style_table(queue_sheet)
    fit_columns(queue_sheet)

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(excel_path)
    except PermissionError:
        alt_path = excel_path.parent / f"{excel_path.stem}_latest{excel_path.suffix}"
        workbook.save(alt_path)
        print(f"[WARN] {excel_path.name} is locked by another program. Saved to {alt_path.name}")
        excel_path = alt_path

    # Re-open once to ensure the generated workbook is structurally readable.
    try:
        checked = load_workbook(excel_path, read_only=True, data_only=False)
        expected = {"Tong quan", "Chi tiet file", "Chi tiet trang", "Can kiem tra", "Review queue"}
        if not expected.issubset(set(checked.sheetnames)):
            raise RuntimeError("Workbook validation failed: missing sheets")
        checked.close()
    except Exception as e:
        print(f"[WARN] Workbook read verification skipped: {e}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export OCR batch results to Excel")
    parser.add_argument("--output-dir", default="output", help="Batch output directory")
    parser.add_argument("--excel", default="output/OCR_BENCHMARK_5_FILES.xlsx", help="Excel output path")
    args = parser.parse_args()
    build_workbook(Path(args.output_dir).resolve(), Path(args.excel).resolve())
    print(f"Excel report created: {Path(args.excel).resolve()}")


if __name__ == "__main__":
    main()
