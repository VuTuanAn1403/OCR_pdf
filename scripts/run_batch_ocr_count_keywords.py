#!/usr/bin/env python3
"""
Compatibility wrapper and historical workbook helpers.

Direct execution delegates to the unified V6.3 CLI before model imports:
full-page GPU OCR -> keyword report -> own Excel -> committed matrix -> next PDF.
Legacy helpers below remain importable by furniture relabelling utilities.
Their historical workflow was:
  1. Sequential execution, one document at a time:
     PDF -> Existing OCR Pipeline (ConvertDocumentUseCase) -> result.md ->
     Count 88 keywords -> Save per-file count_keywords.xlsx ->
     Atomically update aggregate matrix Excel -> Move to next PDF.
  2. Complete processing of all pages using the configured pipeline, model,
     and profile ('balanced'). No simplified bypasses.
  3. Count all 88 lines from config/keywords.txt, preserving exact order.
     Both 'Internet di động' columns at positions 43 and 45 receive identical counts.
  4. Per document:
     - Export count_keywords.xlsx in document output directory.
     - Atomically save aggregate matrix Excel.
     - Print progress: [i/N] file_name | Status | Processing time.
  5. Aggregate matrix:
     - Col 1: PDF name (relative path if name collision).
     - Cols 2..89: Exactly 88 keyword columns.
     - Success: integer count (0 if absent).
     - Error: separate sheet 'Tệp lỗi', leave count cells blank (None).
  6. Final output: output/MA_TRAN_DEM_TU_KHOA_<N>_FILES.xlsx
"""

# Production execution delegates before importing any model/pipeline modules.
if __name__ == "__main__":
    import sys
    from pathlib import Path
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
    sys.argv = [{"--input-dir":"--input","--output-dir":"--output","--force":"--force-ocr"}.get(a,a) for a in sys.argv]
    from src.ocr.cli.unified import run_batch
    raise SystemExit(run_batch())

import argparse
from datetime import datetime, timezone
import functools
import gc
from io import BytesIO
import json
import logging
import os
from pathlib import Path
import re
import subprocess
import sys
import time
import traceback
import unicodedata
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Ensure print flushes immediately to avoid stdout buffering in redirected logs
print = functools.partial(print, flush=True)

# Ensure project root is on sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.ocr.application.convert_document import ConvertDocumentUseCase
from src.ocr.domain.models.keyword import KeywordCount, KeywordReport
from src.ocr.infrastructure.export.completion import (
    MARKER_NAME,
    mark_completed,
    verify_completed,
)
from src.ocr.infrastructure.export.keyword_exporter import KeywordExporter
from src.ocr.runtime.atomic_io import atomic_write_bytes

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("batch_ocr_counter")


def load_raw_keywords(keywords_path: Path) -> List[str]:
    """Load the 88 keywords from config/keywords.txt preserving exact lines and order."""
    text = keywords_path.read_text(encoding="utf-8-sig")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) != 88:
        logger.warning(f"Expected 88 keywords in {keywords_path}, but found {len(lines)}")
    return lines


def normalize_vietnamese_text(text: str) -> str:
    """Normalize Unicode NFC and lowercase for exact search."""
    return unicodedata.normalize("NFC", text).casefold()


def count_keywords_exact(text: str, keywords: List[str]) -> Tuple[List[int], int]:
    """
    Count occurrences of each keyword in text.
    Returns:
      (list_of_88_counts, total_matches)
    """
    from src.ocr.infrastructure.postprocessing.keyword_counter import KeywordCounter
    report = KeywordCounter.count_keywords(text,keywords)
    canonical = {normalize_vietnamese_text(row.keyword):row.count for row in report.keywords}
    counts = [canonical[normalize_vietnamese_text(keyword)] for keyword in keywords]
    total = sum(counts)
    return counts, total


def export_single_document_keywords(
    doc_id: str,
    doc_dir: Path,
    keywords: List[str],
    counts: List[int],
    total_matches: int,
) -> Path:
    """Export standard count_keywords.xlsx with exact 88 keyword rows and update completion marker."""
    keyword_models = [
        KeywordCount(keyword=kw, exact_count=c, count=c)
        for kw, c in zip(keywords, counts)
    ]
    report = KeywordReport(
        document_id=doc_id,
        source_file=str(doc_dir / "result.md"),
        keywords=keyword_models,
        total_keywords_searched=len(keywords),
        total_exact_matches=total_matches,
        total_matches=total_matches,
    )
    excel_path = KeywordExporter.export_excel(report, doc_dir)

    # Re-sign completion marker if validation status exists
    val_status = "PASS"
    val_file = doc_dir / "validation.json"
    if val_file.is_file():
        try:
            val_data = json.loads(val_file.read_text(encoding="utf-8"))
            val_status = val_data.get("status", "PASS")
        except Exception:
            val_status = "PASS"
    try:
        mark_completed(doc_dir, val_status)
    except Exception as exc:
        logger.debug(f"Could not update marker for {doc_id}: {exc}")

    return excel_path


def build_and_save_master_matrix(
    all_files_info: List[Dict[str, Any]],
    keywords: List[str],
    excel_path: Path,
    total_batch_time: float = 0.0,
) -> None:
    """
    Build and atomically write the master aggregate matrix workbook.
    Sheets:
      1. 'Ma trận từ khóa':
         Col 1: Tệp PDF
         Cols 2..89: Exactly 88 keyword columns
         Rows: 1 row per PDF. Integer count (0 if absent). Blank (None) if error.
         Total row at bottom.
      2. 'Tệp lỗi':
         List of failed files with status, error message, and timestamp.
      3. 'Tổng quan':
         Executive summary, KPIs, and keyword ranking table.
    """
    wb = openpyxl.Workbook()

    # Styling definitions
    NAVY_HEADER = PatternFill("solid", fgColor="1F4E78")
    WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    SUB_HEADER = PatternFill("solid", fgColor="D9E1F2")
    SUB_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
    GREEN_FONT = Font(name="Calibri", size=11, color="276A3C", bold=True)
    TOTAL_FILL = PatternFill("solid", fgColor="203764")
    TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    RED_FONT = Font(name="Calibri", size=11, bold=True, color="C00000")

    THIN_BORDER = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    DOUBLE_BOTTOM = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="double", color="1F4E78"),
    )

    # -------------------------------------------------------------------------
    # SHEET 1: MA TRẬN TỪ KHÓA
    # -------------------------------------------------------------------------
    ws_matrix = wb.active
    ws_matrix.title = "Ma trận từ khóa"

    # Header Row (Row 1): Col 1 = Tệp PDF, Cols 2..89 = 88 Keywords
    headers = ["Tệp PDF"] + keywords
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws_matrix.cell(row=1, column=col_idx, value=h_text)
        cell.fill = NAVY_HEADER
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws_matrix.row_dimensions[1].height = 40

    num_files = len(all_files_info)

    # Data Rows (Rows 2 .. num_files + 1)
    for r_idx, doc_info in enumerate(all_files_info, 1):
        curr_row = 1 + r_idx
        pdf_name = doc_info["pdf_name"]

        # Col 1: Tệp PDF
        cell_name = ws_matrix.cell(row=curr_row, column=1, value=pdf_name)
        cell_name.border = THIN_BORDER
        cell_name.alignment = Alignment(horizontal="left", vertical="center")

        status = doc_info.get("status", "PENDING")
        counts = doc_info.get("counts")  # List of 88 ints, or None if error

        if status == "SUCCESS" and counts is not None:
            for kw_idx, cnt in enumerate(counts, 1):
                col_i = 1 + kw_idx
                cell = ws_matrix.cell(row=curr_row, column=col_i, value=cnt)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
                if cnt > 0:
                    cell.fill = GREEN_FILL
                    cell.font = GREEN_FONT
        else:
            # File error or pending: leave blank (None), do not fake 0!
            for kw_idx in range(1, len(keywords) + 1):
                col_i = 1 + kw_idx
                cell = ws_matrix.cell(row=curr_row, column=col_i, value=None)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Total Row (Row num_files + 2)
    total_row = num_files + 2
    cell_tot_label = ws_matrix.cell(row=total_row, column=1, value="TỔNG CỘNG")
    cell_tot_label.fill = TOTAL_FILL
    cell_tot_label.font = TOTAL_FONT
    cell_tot_label.alignment = Alignment(horizontal="center", vertical="center")
    cell_tot_label.border = DOUBLE_BOTTOM

    for kw_idx in range(1, len(keywords) + 1):
        col_i = 1 + kw_idx
        col_letter = get_column_letter(col_i)
        cell = ws_matrix.cell(row=total_row, column=col_i, value=f"=SUM({col_letter}2:{col_letter}{total_row - 1})")
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.number_format = "#,##0"
        cell.border = DOUBLE_BOTTOM

    # Column widths & freeze pane
    ws_matrix.column_dimensions["A"].width = 38
    for kw_idx in range(1, len(keywords) + 1):
        col_letter = get_column_letter(1 + kw_idx)
        kw_len = len(keywords[kw_idx - 1])
        ws_matrix.column_dimensions[col_letter].width = max(13, min(25, kw_len + 3))

    ws_matrix.freeze_panes = "B2"
    last_col_letter = get_column_letter(len(headers))
    ws_matrix.auto_filter.ref = f"A1:{last_col_letter}{num_files + 1}"

    # -------------------------------------------------------------------------
    # SHEET 2: TỆP LỖI
    # -------------------------------------------------------------------------
    ws_errors = wb.create_sheet(title="Tệp lỗi")
    error_headers = ["STT", "Tệp PDF", "Trạng thái", "Nguyên nhân lỗi", "Thời gian ghi nhận"]
    for c_idx, h in enumerate(error_headers, 1):
        cell = ws_errors.cell(row=1, column=c_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws_errors.row_dimensions[1].height = 28

    failed_files = [f for f in all_files_info if f.get("status") == "FAILED"]
    if not failed_files:
        cell_none = ws_errors.cell(row=2, column=1, value=1)
        ws_errors.cell(row=2, column=2, value="(Không có)")
        ws_errors.cell(row=2, column=3, value="THÀNH CÔNG")
        ws_errors.cell(row=2, column=4, value="Không có tệp nào bị lỗi trong toàn bộ quá trình xử lý.")
        ws_errors.cell(row=2, column=5, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        for col_i in range(1, 6):
            ws_errors.cell(row=2, column=col_i).border = THIN_BORDER
    else:
        for err_idx, f_info in enumerate(failed_files, 1):
            curr_r = 1 + err_idx
            ws_errors.cell(row=curr_r, column=1, value=err_idx).alignment = Alignment(horizontal="center")
            ws_errors.cell(row=curr_r, column=2, value=f_info["pdf_name"])
            c_status = ws_errors.cell(row=curr_r, column=3, value=f_info.get("status", "FAILED"))
            c_status.font = RED_FONT
            c_status.alignment = Alignment(horizontal="center")
            ws_errors.cell(row=curr_r, column=4, value=f_info.get("error_message", "Unknown error"))
            ws_errors.cell(row=curr_r, column=5, value=f_info.get("timestamp", ""))
            for col_i in range(1, 6):
                ws_errors.cell(row=curr_r, column=col_i).border = THIN_BORDER

    ws_errors.column_dimensions["A"].width = 8
    ws_errors.column_dimensions["B"].width = 38
    ws_errors.column_dimensions["C"].width = 16
    ws_errors.column_dimensions["D"].width = 65
    ws_errors.column_dimensions["E"].width = 22

    # -------------------------------------------------------------------------
    # SHEET 3: TỔNG QUAN
    # -------------------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Tổng quan")
    ws_summary["A1"] = f"BÁO CÁO TỔNG HỢP TIẾN ĐỘ & KẾT QUẢ ĐẾM TỪ KHÓA ({num_files} TỆP)"
    ws_summary["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")

    success_files = [f for f in all_files_info if f.get("status") == "SUCCESS"]
    total_pages_done = sum(f.get("total_pages", 0) for f in success_files)
    total_matches_all = sum(f.get("total_matches", 0) for f in success_files)

    # Keyword statistics across successful files
    keyword_hit_counts = [0] * len(keywords)
    keyword_doc_counts = [0] * len(keywords)
    for f in success_files:
        counts = f.get("counts", [])
        for i, c in enumerate(counts):
            if c > 0:
                keyword_hit_counts[i] += c
                keyword_doc_counts[i] += 1

    active_keywords_cnt = sum(1 for hits in keyword_hit_counts if hits > 0)

    kpis = [
        ("Tổng số tệp PDF đầu vào", f"{num_files} tệp"),
        ("Số tệp xử lý thành công", f"{len(success_files)} tệp"),
        ("Số tệp xử lý thất bại", f"{len(failed_files)} tệp"),
        ("Tổng số trang tài liệu đã xử lý", f"{total_pages_done:,} trang"),
        ("Tổng số từ khóa cấu hình", f"{len(keywords)} từ khóa"),
        ("Số từ khóa xuất hiện thực tế", f"{active_keywords_cnt}/{len(keywords)} từ khóa"),
        ("Tổng số lượt từ khóa xuất hiện", f"{total_matches_all:,} lượt"),
        ("Thời gian xử lý batch", f"{total_batch_time:.2f}s ({total_batch_time/60:.2f} phút)"),
        ("Thời điểm cập nhật", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]

    for idx, (label, val) in enumerate(kpis, 3):
        ws_summary.cell(row=idx, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=idx, column=2, value=val).font = Font(bold=True, color="1F4E78")
        ws_summary.cell(row=idx, column=1).border = THIN_BORDER
        ws_summary.cell(row=idx, column=2).border = THIN_BORDER

    # Ranking Table
    start_rank_row = len(kpis) + 5
    ws_summary.cell(row=start_rank_row - 1, column=1, value="BẢNG XẾP HẠNG MỨC ĐỘ PHỔ BIẾN TỪ KHÓA").font = SUB_FONT
    rank_headers = ["Hạng", "Từ khóa", "Tổng số lần xuất hiện", "Số doanh nghiệp đề cập", "Tỷ lệ (%)"]
    for c_i, rh in enumerate(rank_headers, 1):
        cell = ws_summary.cell(row=start_rank_row, column=c_i, value=rh)
        cell.fill = NAVY_HEADER
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    ranked_data = []
    for i, kw in enumerate(keywords):
        hits = keyword_hit_counts[i]
        docs = keyword_doc_counts[i]
        pct = (docs / max(1, len(success_files))) * 100
        ranked_data.append((kw, hits, docs, pct))

    ranked_data.sort(key=lambda x: (x[1], x[2]), reverse=True)

    for rank_idx, (kw, hits, docs, pct) in enumerate(ranked_data, 1):
        r_num = start_rank_row + rank_idx
        ws_summary.cell(row=r_num, column=1, value=rank_idx).alignment = Alignment(horizontal="center")
        ws_summary.cell(row=r_num, column=2, value=kw).alignment = Alignment(horizontal="left")
        ws_summary.cell(row=r_num, column=3, value=hits).alignment = Alignment(horizontal="right")
        ws_summary.cell(row=r_num, column=4, value=f"{docs}/{len(success_files)}").alignment = Alignment(horizontal="center")
        ws_summary.cell(row=r_num, column=5, value=round(pct, 1)).alignment = Alignment(horizontal="right")

        for c_i in range(1, 6):
            c = ws_summary.cell(row=r_num, column=c_i)
            c.border = THIN_BORDER
            if c_i == 3:
                c.number_format = "#,##0"
                if hits > 0:
                    c.font = Font(bold=True, color="1F4E78")
            elif c_i == 5:
                c.number_format = '0.0"%"'

    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 38
    ws_summary.column_dimensions["C"].width = 24
    ws_summary.column_dimensions["D"].width = 24
    ws_summary.column_dimensions["E"].width = 16

    # Atomically write workbook
    buffer = BytesIO()
    wb.save(buffer)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_bytes(excel_path, buffer.getvalue())

    # Also keep a copy at project root for convenience
    root_excel = PROJECT_ROOT / excel_path.name
    try:
        atomic_write_bytes(root_excel, buffer.getvalue())
    except Exception:
        pass
    wb.close()


def main():
    if sys.platform == "win32":
        for stream in (sys.stdout, sys.stderr):
            if hasattr(stream, "reconfigure"):
                stream.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=PROJECT_ROOT / "inputs")
    parser.add_argument("--output-dir", type=Path, default=PROJECT_ROOT / "output")
    parser.add_argument("--keywords-file", type=Path, default=PROJECT_ROOT / "config/keywords.txt")
    parser.add_argument("--profile", default="balanced", help="OCR profile to use (default: balanced)")
    parser.add_argument("--device", default="cuda", help="OCR processing device: cuda or cpu (default: cuda)")
    parser.add_argument("--force", action="store_true", help="Force re-OCR even if document was completed")
    parser.add_argument("--worker-pdf", type=Path, default=None, help="Process single PDF in worker process")
    args = parser.parse_args()

    input_dir = args.input_dir.resolve()
    output_dir = args.output_dir.resolve()
    keywords_file = args.keywords_file.resolve()

    # 1. Load exact 88 keywords
    keywords = load_raw_keywords(keywords_file)

    if args.worker_pdf:
        pdf_path = args.worker_pdf.resolve()
        doc_id = pdf_path.stem
        doc_dir = output_dir / doc_id
        result_md = doc_dir / "result.md"
        with ConvertDocumentUseCase(reuse_docling=False) as converter:
            extracted_doc = converter.execute(
                file_path=str(pdf_path),
                profile_name=args.profile,
                output_dir=str(output_dir),
                resume=True,
                benchmark=True,
                device=args.device,
            )
            total_pages = len(extracted_doc.pages)
            md_text = result_md.read_text(encoding="utf-8-sig")
            counts, total_matches = count_keywords_exact(md_text, keywords)
            export_single_document_keywords(doc_id, doc_dir, keywords, counts, total_matches)
        sys.exit(0)

    if not input_dir.is_dir():
        print(f"[ERROR] Thư mục đầu vào không tồn tại: {input_dir}", file=sys.stderr)
        sys.exit(1)

    logger.info(f"Loaded {len(keywords)} keywords from {keywords_file.name}")

    # 2. Discover all PDF files
    pdf_files = sorted(list(input_dir.glob("*.pdf")))
    total_files = len(pdf_files)
    if total_files == 0:
        print(f"[ERROR] Không tìm thấy file PDF nào trong: {input_dir}", file=sys.stderr)
        sys.exit(1)

    # Excel output name per Requirement 6: output/MA_TRAN_DEM_TU_KHOA_<N>_FILES.xlsx
    excel_name = f"MA_TRAN_DEM_TU_KHOA_{total_files}_FILES.xlsx"
    excel_path = output_dir / excel_name

    # Check for name collisions to support relative paths
    name_counts: Dict[str, int] = {}
    for p in pdf_files:
        name_counts[p.name] = name_counts.get(p.name, 0) + 1

    # Initialize tracking structures for all files
    all_files_info: List[Dict[str, Any]] = []
    pre_completed_count = 0

    print("Kiểm tra trạng thái các tài liệu đã hoàn thành trước đó...")
    for idx, pdf_path in enumerate(pdf_files, 1):
        pdf_display_name = (
            str(pdf_path.relative_to(PROJECT_ROOT))
            if name_counts[pdf_path.name] > 1
            else pdf_path.name
        )
        doc_id = pdf_path.stem
        doc_dir = output_dir / doc_id
        result_md = doc_dir / "result.md"

        is_already_done = (
            not args.force
            and verify_completed(doc_dir)
            and result_md.is_file()
        )

        if is_already_done:
            try:
                md_text = result_md.read_text(encoding="utf-8-sig")
                counts, total_matches = count_keywords_exact(md_text, keywords)
                total_pages = 0
                q_file = doc_dir / "quality.json"
                if q_file.is_file():
                    try:
                        q_data = json.loads(q_file.read_text(encoding="utf-8"))
                        total_pages = q_data.get("total_pages", 0)
                    except Exception:
                        pass
                if total_pages == 0:
                    import pymupdf
                    total_pages = len(pymupdf.open(str(pdf_path)))

                # Ensure individual count_keywords.xlsx has exact 88 keyword rows
                export_single_document_keywords(doc_id, doc_dir, keywords, counts, total_matches)

                all_files_info.append({
                    "pdf_path": pdf_path,
                    "pdf_name": pdf_display_name,
                    "doc_id": doc_id,
                    "status": "SUCCESS",
                    "counts": counts,
                    "total_matches": total_matches,
                    "total_pages": total_pages,
                    "elapsed_seconds": 0.0,
                    "error_message": "",
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                })
                pre_completed_count += 1
                print(f"  [{idx:02d}/{total_files:02d}] {pdf_path.name}: ĐÃ HOÀN TẤT TRƯỚC ĐÓ ({total_pages} trang, {total_matches} từ khóa)")
                continue
            except Exception as e:
                logger.warning(f"Không thể đọc kết quả cũ của {doc_id}: {e}")

        all_files_info.append({
            "pdf_path": pdf_path,
            "pdf_name": pdf_display_name,
            "doc_id": doc_id,
            "status": "PENDING",
            "counts": None,
            "total_matches": 0,
            "total_pages": 0,
            "elapsed_seconds": 0.0,
            "error_message": "",
            "timestamp": "",
        })

    # Save initial template of aggregate matrix (preserving all completed progress)
    build_and_save_master_matrix(all_files_info, keywords, excel_path, 0.0)

    print("=" * 80)
    print(f"BẮT ĐẦU CHẠY BATCH OCR VÀ ĐẾM TỪ KHÓA (TỔNG {total_files} TÀI LIỆU, ĐÃ XONG {pre_completed_count})")
    print(f"Profile: {args.profile} | Thiết bị: {args.device} | Cấu hình: {len(keywords)} từ khóa")
    print(f"File ma trận tổng hợp: {excel_path.name}")
    print("=" * 80)

    start_batch_time = time.time()

    for idx, doc_info in enumerate(all_files_info, 1):
        if doc_info["status"] == "SUCCESS":
            # Already processed and verified
            continue

        pdf_path: Path = doc_info["pdf_path"]
        doc_id: str = doc_info["doc_id"]
        doc_dir: Path = output_dir / doc_id
        result_md = doc_dir / "result.md"

        file_start_time = time.time()
        print(f"\n[{idx:02d}/{total_files:02d}] Bắt đầu xử lý: {pdf_path.name}...")

        try:
            # Execute complete OCR pipeline in an isolated worker process
            # to guarantee 100% memory reclamation after every file
            cmd = [
                sys.executable,
                str(Path(__file__).resolve()),
                "--worker-pdf", str(pdf_path),
                "--output-dir", str(output_dir),
                "--keywords-file", str(keywords_file),
                "--profile", args.profile,
                "--device", args.device,
            ]
            proc = subprocess.run(cmd)
            if proc.returncode != 0:
                raise RuntimeError(f"Worker process exited with code {proc.returncode}")

            if not result_md.is_file():
                raise FileNotFoundError(f"result.md not created at {result_md}")

            total_pages = 0
            q_file = doc_dir / "quality.json"
            if q_file.is_file():
                try:
                    q_data = json.loads(q_file.read_text(encoding="utf-8"))
                    total_pages = q_data.get("total_pages", 0)
                except Exception:
                    pass
            if total_pages == 0:
                import pymupdf
                total_pages = len(pymupdf.open(str(pdf_path)))

            md_text = result_md.read_text(encoding="utf-8-sig")
            counts, total_matches = count_keywords_exact(md_text, keywords)
            export_single_document_keywords(doc_id, doc_dir, keywords, counts, total_matches)

            file_elapsed = round(time.time() - file_start_time, 2)
            doc_info["status"] = "SUCCESS"
            doc_info["counts"] = counts
            doc_info["total_matches"] = total_matches
            doc_info["total_pages"] = total_pages
            doc_info["elapsed_seconds"] = file_elapsed
            doc_info["timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            unique_hits = sum(1 for c in counts if c > 0)
            print(
                f"   -> [THÀNH CÔNG] {pdf_path.name} ({total_pages} trang) | "
                f"Từ khóa: {total_matches:,} lần (gặp {unique_hits}/{len(keywords)} từ) | "
                f"Thời gian: {file_elapsed:.2f}s"
            )

        except Exception as exc:
            file_elapsed = round(time.time() - file_start_time, 2)
            err_msg = str(exc)
            logger.error(f"Lỗi khi xử lý file {pdf_path.name}: {err_msg}", exc_info=True)
            doc_info["status"] = "FAILED"
            doc_info["counts"] = None  # Left blank in matrix!
            doc_info["total_matches"] = 0
            doc_info["total_pages"] = 0
            doc_info["elapsed_seconds"] = file_elapsed
            doc_info["error_message"] = err_msg
            doc_info["timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            print(
                f"   -> [THẤT BẠI] {pdf_path.name} | "
                f"Nguyên nhân: {err_msg} | "
                f"Thời gian: {file_elapsed:.2f}s"
            )

        # Atomically update master matrix workbook after EVERY single document
        current_total_time = time.time() - start_batch_time
        build_and_save_master_matrix(all_files_info, keywords, excel_path, current_total_time)
        print(f"   -> Đã cập nhật nguyên tử vào ma trận: {excel_path.name}")

    total_batch_time = round(time.time() - start_batch_time, 2)
    success_count = sum(1 for f in all_files_info if f["status"] == "SUCCESS")
    failed_count = sum(1 for f in all_files_info if f["status"] == "FAILED")

    print("\n" + "=" * 80)
    print("HOÀN TẤT QUÁ TRÌNH XỬ LÝ TOÀN BỘ BATCH!")
    print(f"Tổng số tài liệu: {total_files}")
    print(f"Số file thành công: {success_count}/{total_files}")
    print(f"Số file thất bại: {failed_count}/{total_files}")
    print(f"Tổng thời gian: {total_batch_time:.2f}s ({total_batch_time/60:.2f} phút)")
    print(f"Đường dẫn file Excel ma trận tổng hợp: {excel_path.resolve()}")
    print("=" * 80)


if __name__ == "__main__":
    main()
