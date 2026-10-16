"""V6.3 workbook projections. All totals are values, never cached formulas."""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from src.ocr.runtime.atomic_io import atomic_write_bytes
from src.ocr.domain.models.keyword import KeywordReport, KeywordCount
from src.ocr.infrastructure.export.keyword_exporter import KeywordExporter


def export_document(report, snapshot, output):
    rows = [KeywordCount(keyword=k, count=n, exact_count=n) for k, n in zip(snapshot["display_columns"], report["display_counts"])]
    model = KeywordReport(document_id=report["document_id"], source_file=report["source_file"], keywords=rows,
                          total_keywords_searched=len(rows), total_matches=report["total_display_matches"],
                          total_exact_matches=report["total_display_matches"])
    path = KeywordExporter.export_excel(model, Path(output))
    # Include both totals and the policy without changing the established rows.
    book = load_workbook(path)
    try:
        book.active["A3"] = (f"{snapshot['keyword_set_version']} | literal_v1 | canonical={report['total_canonical_matches']} | "
                             f"display={report['total_display_matches']} | {snapshot['snapshot_sha256']}")
        buffer = BytesIO()
        book.save(buffer)
        atomic_write_bytes(path, buffer.getvalue())
    finally:
        book.close()
    verify_document(path, report, snapshot)


def verify_document(path, report, snapshot):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book.active.iter_rows(min_row=5, max_row=4+len(snapshot["display_columns"]), values_only=True))
        if [r[1] for r in rows] != snapshot["display_columns"] or [r[2] for r in rows] != report["display_counts"]:
            raise ValueError("Per-document Excel differs from canonical keyword report")
    finally:
        book.close()


def matrix_bytes(batch_id, revision, records, snapshot):
    book = Workbook(write_only=True)
    try:
        sheet = book.create_sheet("Ma trận từ khóa")
        header = ["Tệp PDF", *snapshot["display_columns"]]
        # Write-only cells explicitly treat dictionary entries as literal text.
        from openpyxl.cell import WriteOnlyCell
        cells = [WriteOnlyCell(sheet, value=s) for s in header]
        for cell in cells:
            cell.data_type = "s"
        sheet.append(cells)
        totals = [0]*len(snapshot["display_columns"])
        for row in records:
            counts = row.get("counts")
            name = WriteOnlyCell(sheet, value=row["name"])
            name.data_type = "s"
            sheet.append([name, *(counts if counts is not None else [None]*len(totals))])
            if counts is not None:
                totals = [a+b for a, b in zip(totals, counts)]
        sheet.append(["TỔNG CỘNG", *totals])
        status = book.create_sheet("Trạng thái")
        status.append(["Document ID", "Tệp PDF", "Kết quả", "Validation", "Attempt", "Report hash", "Publication", "Revision", "Lỗi", "Source review", "Pages", "Duplicate of"])
        errors = book.create_sheet("Tệp lỗi")
        errors.append(["Document ID", "Lỗi"])
        for row in records:
            status.append([row["id"], row["name"], row.get("outcome", "PENDING"), row.get("validation"), row.get("attempt", 0),
                           row.get("report_sha256"), "INCLUDED_IN_REVISION" if row.get("outcome") else "PENDING", revision, row.get("error"),row.get("source_review"),row.get("total_pages"),row.get("duplicate_of")])
            if row.get("outcome") == "FAILED_DOCUMENT":
                errors.append([row["id"], row.get("error")])
        summary = book.create_sheet("Tổng quan")
        for key, value in {"Tổng đầu vào": len(records), "Đã công bố": sum(bool(r.get("outcome")) for r in records),
            "Đang chờ": sum(not r.get("outcome") for r in records), "Lỗi": sum(r.get("outcome")=="FAILED_DOCUMENT" for r in records),
            "Cần review": sum(r.get("validation")=="NEEDS_REVIEW" for r in records),
            "Nguồn hash duy nhất": len({r["pdf_sha256"] for r in records if r.get("pdf_sha256")}),
            "Tổng canonical": sum(r.get("total_canonical_matches", 0) for r in records),
            "Tổng display": sum(totals)}.items():
            summary.append([key, value])
        info = book.create_sheet("Thông tin")
        for key, value in {"batch_id": batch_id, "revision": revision, "dictionary_sha256": snapshot["snapshot_sha256"],
                           "schema_version": "keyword-matrix-v1", "policy": "literal_v1", "quality_acceptance": "NOT_EVALUATED"}.items():
            info.append([key, value])
        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue()
    finally:
        book.close()


def verify_matrix(path, batch_id, revision, records, snapshot):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        info = dict(book["Thông tin"].iter_rows(values_only=True))
        if info.get("batch_id") != batch_id or info.get("revision") != revision or info.get("dictionary_sha256") != snapshot["snapshot_sha256"]:
            raise ValueError("Matrix identity/revision mismatch")
        rows = iter(book["Ma trận từ khóa"].iter_rows(min_col=1,max_col=1+len(snapshot["display_columns"]),values_only=True))
        if list(next(rows)) != ["Tệp PDF", *snapshot["display_columns"]]:
            raise ValueError("Matrix header mismatch")
        totals = [0]*len(snapshot["display_columns"])
        for record in records:
            values = list(next(rows))
            expected = record.get("counts")
            if values != [record["name"], *(expected if expected is not None else [None]*len(totals))]:
                raise ValueError("Matrix rows/counts mismatch")
            if expected is not None:
                totals = [a+b for a, b in zip(totals, expected)]
        if list(next(rows)) != ["TỔNG CỘNG", *totals] or next(rows, None) is not None:
            raise ValueError("Matrix totals/row count mismatch")
    finally:
        book.close()
