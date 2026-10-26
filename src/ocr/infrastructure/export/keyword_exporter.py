"""Excel keyword report with two sheets inside a single count_keywords.xlsx.

Sheet 1 – ``Tong_hop``
    Flat feature-matrix: one row per document, one column per keyword.

Sheet 2 – ``Trich_dan_chi_tiet``
    Per-occurrence citation table with page numbers and ±10-word context
    windows.  The keyword inside each citation is rendered in **bold**
    using Excel rich-text formatting (openpyxl CellRichText).
"""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.ocr.domain.models.keyword import KeywordReport
from src.ocr.runtime.atomic_io import atomic_write_bytes

# ── Try to import rich-text support (openpyxl ≥ 3.1) ──
try:
    from openpyxl.cell.rich_text import TextBlock, CellRichText
    from openpyxl.cell.text import InlineFont
    _HAS_RICH_TEXT = True
except ImportError:
    _HAS_RICH_TEXT = False

# ── Shared style constants ──
_NAVY = "1F4E79"
_HEADER_FILL = PatternFill("solid", fgColor=_NAVY)
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FONT_SM = Font(bold=True, color="FFFFFF", size=9)
_CENTER = Alignment(horizontal="center", vertical="center")
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_BOLD = Font(bold=True)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class KeywordExporter:
    @staticmethod
    def export_excel(report: KeywordReport, output_dir: Path) -> Path:
        workbook = Workbook()
        try:
            _build_summary_sheet(workbook, report)
            _build_citation_sheet(workbook, report)
            buffer = BytesIO()
            workbook.save(buffer)
            path = Path(output_dir) / "count_keywords.xlsx"
            atomic_write_bytes(path, buffer.getvalue())
            return path
        finally:
            workbook.close()

    @staticmethod
    def export_file_excel(report: KeywordReport, output_file: Path) -> Path:
        workbook = Workbook()
        try:
            _build_summary_sheet(workbook, report)
            _build_citation_sheet(workbook, report)
            buffer = BytesIO()
            workbook.save(buffer)
            path = Path(output_file)
            path.parent.mkdir(parents=True, exist_ok=True)
            atomic_write_bytes(path, buffer.getvalue())
            return path
        finally:
            workbook.close()

    @staticmethod
    def export_batch_excel(reports: list[KeywordReport], output_file: Path) -> Path:
        workbook = Workbook()
        try:
            _build_batch_summary_sheet(workbook, reports)
            _build_batch_citation_sheet(workbook, reports)
            buffer = BytesIO()
            workbook.save(buffer)
            path = Path(output_file)
            path.parent.mkdir(parents=True, exist_ok=True)
            atomic_write_bytes(path, buffer.getvalue())
            return path
        finally:
            workbook.close()


# ─────────────────────────────────────────────────────────────────────
# Sheet 1: Tong_hop  (Feature-matrix)
# ─────────────────────────────────────────────────────────────────────

def _build_summary_sheet(workbook: Workbook, report: KeywordReport):
    ws = workbook.active
    ws.title = "Tong_hop"

    # ── Header row ──
    file_cell = ws.cell(1, 1, "File_name")
    file_cell.font = _HEADER_FONT
    file_cell.fill = _HEADER_FILL
    file_cell.alignment = _CENTER
    file_cell.border = _THIN_BORDER

    for col_idx, kw in enumerate(report.keywords, 2):
        cell = ws.cell(1, col_idx, kw.keyword)
        cell.font = _HEADER_FONT_SM
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER_WRAP
        cell.border = _THIN_BORDER

    # ── Data row ──
    name_cell = ws.cell(2, 1, report.document_id)
    name_cell.border = _THIN_BORDER

    for col_idx, kw in enumerate(report.keywords, 2):
        cell = ws.cell(2, col_idx, kw.count)
        cell.number_format = "#,##0"
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    # ── Column widths ──
    ws.column_dimensions["A"].width = max(24, len(report.document_id) + 4)
    for col_idx in range(2, len(report.keywords) + 2):
        col_letter = get_column_letter(col_idx)
        kw_len = len(report.keywords[col_idx - 2].keyword)
        ws.column_dimensions[col_letter].width = max(8, min(22, kw_len + 2))

    ws.freeze_panes = "B2"


# ─────────────────────────────────────────────────────────────────────
# Sheet 2: Trich_dan_chi_tiet  (Per-occurrence citations)
# ─────────────────────────────────────────────────────────────────────

def _build_citation_sheet(workbook: Workbook, report: KeywordReport):
    ws = workbook.create_sheet("Trich_dan_chi_tiet")

    # ── Header row ──
    headers = ["STT", "Tên từ khóa", "Thuộc trang số", "Trích dẫn"]
    col_widths = [8, 32, 16, 90]
    for col_idx, (title, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(1, col_idx, title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    if not report.citations:
        ws.freeze_panes = "A2"
        return

    # ── Sort: keyword order in keywords.txt, then page number ──
    keyword_order = {kw.keyword.casefold(): i for i, kw in enumerate(report.keywords)}
    sorted_citations = sorted(
        report.citations,
        key=lambda c: (keyword_order.get(c.keyword.casefold(), 9999), c.page_number),
    )

    # ── Auto-fit column B to longest keyword ──
    max_kw_len = max((len(c.keyword) for c in sorted_citations), default=20)
    ws.column_dimensions["B"].width = max(24, min(44, max_kw_len + 6))

    # ── Data rows ──
    for row_idx, cit in enumerate(sorted_citations, 1):
        row = row_idx + 1

        # STT
        stt_cell = ws.cell(row, 1, row_idx)
        stt_cell.alignment = _CENTER
        stt_cell.border = _THIN_BORDER

        # Tên từ khóa
        kw_cell = ws.cell(row, 2, cit.keyword)
        kw_cell.border = _THIN_BORDER

        # Thuộc trang số
        page_cell = ws.cell(row, 3, cit.page_number)
        page_cell.alignment = _CENTER
        page_cell.border = _THIN_BORDER

        # Trích dẫn (with bold keyword via rich text)
        cite_cell = ws.cell(row, 4)
        cite_cell.alignment = _WRAP
        cite_cell.border = _THIN_BORDER
        _write_citation_cell(cite_cell, cit)

    ws.freeze_panes = "A2"


def _write_citation_cell(cell, citation) -> None:
    """Write the citation with the keyword portion in bold.

    Uses openpyxl CellRichText when available; otherwise falls back to
    plain text with bracket markers ``[keyword]``.
    """
    before = citation.citation_before
    kw = citation.citation_keyword
    after = citation.citation_after

    if _HAS_RICH_TEXT:
        parts: list[TextBlock] = []
        bold_font = InlineFont(b=True)
        normal_font = InlineFont()
        if before:
            parts.append(TextBlock(normal_font, before + " "))
        parts.append(TextBlock(bold_font, kw))
        if after:
            parts.append(TextBlock(normal_font, " " + after))
        if parts:
            cell.value = CellRichText(*parts)
    else:
        # Fallback: bracket-mark the keyword
        segments = []
        if before:
            segments.append(before)
        segments.append(f"[{kw}]")
        if after:
            segments.append(after)
        cell.value = " ".join(segments)


# ─────────────────────────────────────────────────────────────────────
# Batch Sheets: Tong_hop & Trich_dan_chi_tiet across multiple reports
# ─────────────────────────────────────────────────────────────────────

def _build_batch_summary_sheet(workbook: Workbook, reports: list[KeywordReport]):
    ws = workbook.active
    ws.title = "Tong_hop"
    if not reports:
        return

    keywords = reports[0].keywords

    # ── Header row ──
    file_cell = ws.cell(1, 1, "File_name")
    file_cell.font = _HEADER_FONT
    file_cell.fill = _HEADER_FILL
    file_cell.alignment = _CENTER
    file_cell.border = _THIN_BORDER

    for col_idx, kw in enumerate(keywords, 2):
        cell = ws.cell(1, col_idx, kw.keyword)
        cell.font = _HEADER_FONT_SM
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER_WRAP
        cell.border = _THIN_BORDER

    # ── Data rows ──
    for row_idx, rep in enumerate(reports, 2):
        name_cell = ws.cell(row_idx, 1, rep.document_id)
        name_cell.border = _THIN_BORDER

        for col_idx, kw in enumerate(rep.keywords, 2):
            cell = ws.cell(row_idx, col_idx, kw.count)
            cell.number_format = "#,##0"
            cell.alignment = _CENTER
            cell.border = _THIN_BORDER

    # ── Total row ──
    total_row = len(reports) + 2
    tot_label = ws.cell(total_row, 1, "TỔNG CỘNG")
    tot_label.font = _BOLD
    tot_label.alignment = _CENTER
    tot_label.border = _THIN_BORDER

    for col_idx in range(2, len(keywords) + 2):
        col_letter = get_column_letter(col_idx)
        cell = ws.cell(total_row, col_idx, f"=SUM({col_letter}2:{col_letter}{total_row - 1})")
        cell.font = _BOLD
        cell.number_format = "#,##0"
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    # ── Column widths ──
    max_name_len = max((len(r.document_id) for r in reports), default=20)
    ws.column_dimensions["A"].width = max(24, max_name_len + 4)
    for col_idx in range(2, len(keywords) + 2):
        col_letter = get_column_letter(col_idx)
        kw_len = len(keywords[col_idx - 2].keyword)
        ws.column_dimensions[col_letter].width = max(8, min(22, kw_len + 2))

    ws.freeze_panes = "B2"


def _build_batch_citation_sheet(workbook: Workbook, reports: list[KeywordReport]):
    ws = workbook.create_sheet("Trich_dan_chi_tiet")

    # ── Header row ──
    headers = ["STT", "Tên file", "Tên từ khóa", "Thuộc trang số", "Trích dẫn"]
    col_widths = [8, 30, 28, 16, 90]
    for col_idx, (title, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(1, col_idx, title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    all_cits: list[tuple[str, int, any]] = []
    for rep in reports:
        keyword_order = {kw.keyword.casefold(): i for i, kw in enumerate(rep.keywords)}
        for c in rep.citations:
            all_cits.append((rep.document_id, keyword_order.get(c.keyword.casefold(), 9999), c))

    if not all_cits:
        ws.freeze_panes = "A2"
        return

    # Sort by Document ID, then keyword order in keywords.txt, then page number
    all_cits.sort(key=lambda item: (item[0], item[1], item[2].page_number))

    for row_idx, (doc_id, _, cit) in enumerate(all_cits, 1):
        row = row_idx + 1

        # STT
        ws.cell(row, 1, row_idx).alignment = _CENTER
        ws.cell(row, 1).border = _THIN_BORDER

        # Tên file
        ws.cell(row, 2, doc_id).border = _THIN_BORDER

        # Tên từ khóa
        ws.cell(row, 3, cit.keyword).border = _THIN_BORDER

        # Thuộc trang số
        page_cell = ws.cell(row, 4, cit.page_number)
        page_cell.alignment = _CENTER
        page_cell.border = _THIN_BORDER

        # Trích dẫn (with bold keyword via rich text)
        cite_cell = ws.cell(row, 5)
        cite_cell.alignment = _WRAP
        cite_cell.border = _THIN_BORDER
        _write_citation_cell(cite_cell, cit)

    ws.freeze_panes = "A2"

